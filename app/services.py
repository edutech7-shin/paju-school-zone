from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
import math
import numbers
import os
import tempfile
import re
import threading
import time
from typing import Any
from uuid import uuid4

import pandas as pd
from openpyxl import load_workbook
import requests

from addressConverter import get_jibeon_address
from mapping_final import (
    extract_address_components,
    extract_building_name,
    parse_district_ranges,
    map_address_to_school_area,
    manual_mapping,
    standardize_address,
)

from .config import load_admin_config
from .datasets import ensure_datasets
from .utils import any_contains, clean_text, compact_text, parcel_in_ranges, parse_jibun_address, tokenize_korean_terms
from .utils import parse_address_unit_info, unit_in_ranges


GRADE_COLUMNS = [f"{grade}학년" for grade in range(1, 7)]


@dataclass
class AppContext:
    admin_records: list[dict[str, Any]]
    report_template_records: list[dict[str, Any]]
    school_zone_records: list[dict[str, Any]]
    fallback_district_ranges: pd.DataFrame


_CONTEXT_CACHE: dict[tuple[str, str, str, str], AppContext] = {}
_CLEANUP_THREAD_STARTED = False

ADDRESS_ANCHORS: list[dict[str, str]] = [
    {
        "id": "yadang-1015-hanbit4",
        "contains": "야당동1015",
        "contains_any": "롯데캐슬파크타운2차|롯데캐슬파크타운ii|롯데캐슬파크타운ⅱ|롯데캐슬파크타운Ⅱ|한빛마을4단지",
        "admin_area": "운정3동",
        "tong_ri": "31통",
        "school_name": "와석초",
    },
]


def get_context() -> AppContext:
    config = load_admin_config()
    key = (
        config.data_sources.admin_source,
        config.data_sources.report_template_source,
        config.data_sources.school_zone_source,
        config.data_sources.legacy_district_source,
    )
    if key in _CONTEXT_CACHE:
        return _CONTEXT_CACHE[key]
    datasets = ensure_datasets()
    fallback_district_ranges = _load_fallback_district_ranges(config.data_sources.legacy_district_source)
    context = AppContext(
        admin_records=datasets["admin"],
        report_template_records=datasets["report_template"],
        school_zone_records=datasets["school_zones"],
        fallback_district_ranges=fallback_district_ranges,
    )
    _CONTEXT_CACHE.clear()
    _CONTEXT_CACHE[key] = context
    return context


def _load_fallback_district_ranges(source_path: str) -> pd.DataFrame:
    """
    Vercel 배포 환경에서는 로컬 절대경로 원본 엑셀이 없을 수 있으므로,
    파일이 없거나 읽기에 실패해도 빈 범위 테이블로 안전하게 처리합니다.
    """
    empty_columns = ["동이름", "관할동", "관할통", "시작본번", "시작부번", "끝본번", "끝부번"]
    if not source_path or not os.path.exists(source_path):
        return pd.DataFrame(columns=empty_columns)
    try:
        district_df = pd.read_excel(source_path)
        parsed = parse_district_ranges(district_df)
        if parsed is None or parsed.empty:
            return pd.DataFrame(columns=empty_columns)
        return parsed
    except Exception:
        return pd.DataFrame(columns=empty_columns)


def _resolve_report_template_path(source_path: str) -> str:
    """
    1) 로컬 경로가 있으면 우선 사용
    2) 없으면 REPORT_TEMPLATE_URL에서 /tmp 캐시에 1회 다운로드 후 사용
    """
    if source_path and os.path.exists(source_path):
        return source_path

    template_url = os.getenv("REPORT_TEMPLATE_URL", "").strip()
    if not template_url:
        return ""

    cache_dir = os.path.join(tempfile.gettempdir(), "addressConvert", "templates")
    os.makedirs(cache_dir, exist_ok=True)
    cached_template_path = os.path.join(cache_dir, "report_template.xlsx")
    if os.path.exists(cached_template_path):
        return cached_template_path

    response = requests.get(template_url, timeout=30)
    response.raise_for_status()
    with open(cached_template_path, "wb") as file:
        file.write(response.content)
    return cached_template_path


def clear_context_cache() -> None:
    _CONTEXT_CACHE.clear()


def clear_runtime_data() -> dict[str, int]:
    with JOB_LOCK:
        result_count = len(RESULT_DATA_STORE)
        file_count = len(RESULT_STORE)
        job_count = len(JOB_STORE)
        RESULT_DATA_STORE.clear()
        RESULT_STORE.clear()
        JOB_STORE.clear()
    return {"results": result_count, "downloads": file_count, "jobs": job_count}


def _safe_int(value, default=0) -> int:
    try:
        if pd.isna(value):
            return default
    except TypeError:
        pass
    text = str(value).strip()
    if not text:
        return default
    match = re.search(r"\d+", text)
    return int(match.group(0)) if match else default


def _format_base_date(base_date: str) -> str:
    text = clean_text(base_date)
    if not text:
        return ""
    try:
        parsed = datetime.strptime(text, "%Y-%m-%d")
        return f"{parsed.year}. {parsed.month}. {parsed.day}."
    except ValueError:
        return text


def _resolve_address_anchor(
    road_address: str,
    jibun_address: str,
    building_name: str,
    requested_school_name: str = "",
) -> dict[str, str] | None:
    joined = compact_text(" ".join(filter(None, [road_address, jibun_address, building_name])))
    if not joined:
        return None
    for rule in ADDRESS_ANCHORS:
        must = compact_text(rule.get("contains", ""))
        if must and must not in joined:
            continue
        any_tokens = [compact_text(token) for token in clean_text(rule.get("contains_any", "")).split("|") if token]
        if any_tokens and not any(token in joined for token in any_tokens):
            continue
        anchor_school = clean_text(rule.get("school_name", ""))
        if requested_school_name and anchor_school and requested_school_name != anchor_school:
            continue
        return rule
    return None


def _is_outside_paju_address(road_address: str, jibun_address: str) -> bool:
    text = clean_text(" ".join(filter(None, [road_address, jibun_address])))
    if not text:
        return False
    compact = compact_text(text)
    if "경기도파주시" in compact:
        return False
    # 시/도 정보가 있는데 파주시가 아닌 경우는 관외로 본다.
    return bool(
        re.search(
            r"(특별시|광역시|특별자치시|특별자치도|경기도|강원도|충청북도|충청남도|전라북도|전라남도|경상북도|경상남도|제주특별자치도)",
            text,
        )
    )


def _find_outside_special_row(context: AppContext, school_name: str) -> dict[str, Any]:
    for record in context.report_template_records:
        if (
            record.get("sheet_name") == "재학생 현황(관내)"
            and record.get("category") == "special"
            and record.get("admin_area") == "기타(통학구역외)"
            and (not school_name or record.get("school_name") == school_name)
        ):
            return record
    return {}


def _find_column(df: pd.DataFrame, candidates: list[str]) -> str | None:
    normalized = {compact_text(col): col for col in df.columns}
    for candidate in candidates:
        key = compact_text(candidate)
        if key in normalized:
            return normalized[key]
    return None


def _excel_cell_display(value: object) -> str:
    """엑셀 셀 값을 명렬표에 적힌 것처럼 문자열로 유지한다(반·번호 등)."""
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except TypeError:
        pass
    if isinstance(value, bool):
        return clean_text(str(value))
    if isinstance(value, numbers.Integral):
        return str(int(value))
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        if value.is_integer():
            return str(int(value))
    return clean_text(str(value))


def _extract_student_rows(df: pd.DataFrame, fallback_grade: int | None = None) -> list[dict[str, Any]]:
    grade_col = _find_column(df, ["학년"])
    name_col = _find_column(df, ["성명", "이름", "학생명"])
    zip_col = _find_column(df, ["우편번호", "우편 번호", "zip", "zipcode"])
    address_col = _find_column(df, ["주소", "도로명주소", "주소지"])
    jibun_col = _find_column(df, ["지번주소"])
    homeroom_col = _find_column(df, ["반", "학급"])
    number_col = _find_column(df, ["번호", "학생번호", "출석번호", "no"])

    if address_col is None and jibun_col is None:
        raise ValueError("학생 파일에 '주소' 또는 '지번주소' 컬럼이 필요합니다.")

    rows = []
    for _, raw in df.iterrows():
        grade = _safe_int(raw.get(grade_col) if grade_col else fallback_grade, fallback_grade or 0)
        name = clean_text(raw.get(name_col) if name_col else "")
        postal_code = clean_text(raw.get(zip_col) if zip_col else "")
        road_address = clean_text(raw.get(address_col) if address_col else "")
        jibun_address = clean_text(raw.get(jibun_col) if jibun_col else "")
        homeroom_ban = _excel_cell_display(raw.get(homeroom_col)) if homeroom_col else ""
        attendance_no = _excel_cell_display(raw.get(number_col)) if number_col else ""

        if not road_address and not jibun_address:
            continue

        rows.append(
            {
                "grade": grade,
                "name": name,
                "postal_code": postal_code,
                "road_address": road_address,
                "source_jibun_address": jibun_address,
                "homeroom_ban": homeroom_ban,
                "attendance_no": attendance_no,
            }
        )
    return rows


def load_student_files(files: list[tuple[str, bytes]]) -> list[dict[str, Any]]:
    all_rows: list[dict[str, Any]] = []
    for filename, payload in files:
        xls = pd.ExcelFile(BytesIO(payload))
        fallback_grade = _safe_int(filename, 0) or None
        file_rows = []
        for sheet_name in xls.sheet_names:
            df = pd.read_excel(BytesIO(payload), sheet_name=sheet_name)
            file_rows.extend(_extract_student_rows(df, fallback_grade=fallback_grade))
        all_rows.extend(file_rows)
    if not all_rows:
        raise ValueError("학생 주소 데이터를 읽지 못했습니다.")
    return all_rows


def _extract_tong_number(text: str) -> str:
    match = re.search(r"(\d+(?:~\d+)?)통", clean_text(text))
    return match.group(1) if match else ""


def _normalize_road_for_jibun_lookup(road_address: str) -> str:
    text = clean_text(road_address)
    if not text:
        return ""
    # "601동 1903"처럼 동만 있고 호 접미사가 누락된 경우, 호를 자동 보정한다.
    text = re.sub(r"(\d+동)\s+(\d+)\s*$", r"\1 \2호", text)
    # "509-302호", "612-901호" 같은 동-호 표현은 API 검색 정확도를 떨어뜨려 제거한다.
    text = re.sub(r"\s*\d{3,4}-\d{2,4}호?\s*$", "", text)
    # 쉼표 뒤 상세호수 정보는 우선 제거 후 지번 변환을 시도한다.
    text = re.sub(r"\s*,\s*\d+동\s*\d+호\s*$", "", text)
    return text.strip()


def _extract_tong_numbers(text: str) -> set[int]:
    numbers: set[int] = set()
    for start_text, end_text in re.findall(r"(\d+)(?:~(\d+))?통", clean_text(text)):
        start = int(start_text)
        end = int(end_text or start_text)
        if start > end:
            start, end = end, start
        numbers.update(range(start, end + 1))
    return numbers


def _extract_ri_numbers(text: str) -> dict[str, set[int]]:
    """
    "용미1~4리, 분수2~3리" 같은 문자열에서 리 이름별 번호 집합을 만든다.
    """
    values: dict[str, set[int]] = {}
    normalized = clean_text(text)
    for name, start_text, end_text in re.findall(r"([가-힣]+)\s*(\d+)(?:~(\d+))?리", normalized):
        start = int(start_text)
        end = int(end_text or start_text)
        if start > end:
            start, end = end, start
        bucket = values.setdefault(name, set())
        bucket.update(range(start, end + 1))
    return values


def _extract_ri_base_names(text: str) -> set[str]:
    names: set[str] = set()
    normalized = clean_text(text)
    # "용미1~4리", "분수2리", "용미리" 형태 모두에서 기본 리 이름을 추출한다.
    for name in re.findall(r"([가-힣]+)\s*\d+(?:~\d+)?리", normalized):
        if name:
            names.add(name)
    for name in re.findall(r"([가-힣]+)리", normalized):
        if name:
            names.add(name)
    return names


def _tong_matches(student_tong: str, candidate_tong: str) -> bool:
    student_numbers = _extract_tong_numbers(student_tong)
    candidate_numbers = _extract_tong_numbers(candidate_tong)
    if not student_numbers or not candidate_numbers:
        return False
    return bool(student_numbers & candidate_numbers)


def _tong_or_ri_matches(student_tong: str, candidate_tong: str) -> bool:
    # 기존 통(通) 비교
    if _tong_matches(student_tong, candidate_tong):
        return True

    # 리(里) 비교: 같은 리 이름 + 번호 교집합이 있으면 일치로 본다.
    student_ri = _extract_ri_numbers(student_tong)
    candidate_ri = _extract_ri_numbers(candidate_tong)
    if student_ri and candidate_ri:
        for name, student_numbers in student_ri.items():
            if name in candidate_ri and (student_numbers & candidate_ri[name]):
                return True

    # "용미리" vs "용미1~4리"처럼 번호가 없는 리 이름만 있어도 같은 리면 일치 처리
    student_ri_bases = _extract_ri_base_names(student_tong)
    candidate_ri_bases = _extract_ri_base_names(candidate_tong)
    if student_ri_bases and candidate_ri_bases and (student_ri_bases & candidate_ri_bases):
        return True

    # 범위 파싱이 불가능한 경우를 위한 최소 문자열 포함 비교
    # 단, 통(通)끼리는 "18통"에 "8통"이 부분문자열로 포함되는 오탐이 있으므로 적용하지 않는다.
    student_compact = compact_text(student_tong)
    candidate_compact = compact_text(candidate_tong)
    if (
        student_compact
        and candidate_compact
        and ("리" in student_tong or "리" in candidate_tong)
    ):
        if student_compact in candidate_compact or candidate_compact in student_compact:
            return True
    return False


def _parse_school_names(text: str) -> set[str]:
    return {token.strip() for token in re.split(r"[,/\n]+", clean_text(text)) if token.strip().endswith("초")}


def _is_joint_zone_pair(requested_row: dict[str, Any], actual_row: dict[str, Any], requested_school_name: str, actual_school_name: str) -> bool:
    if not requested_row.get("row_number") or not actual_row.get("row_number"):
        return False
    if requested_row.get("admin_area") != actual_row.get("admin_area"):
        return False
    if not _tong_or_ri_matches(requested_row.get("tong_ri", ""), actual_row.get("tong_ri", "")):
        return False
    if compact_text(requested_row.get("extra_area", "")) != compact_text(actual_row.get("extra_area", "")):
        return False
    requested_note_schools = _parse_school_names(requested_row.get("note", ""))
    actual_note_schools = _parse_school_names(actual_row.get("note", ""))
    return actual_school_name in requested_note_schools or requested_school_name in actual_note_schools


def _rule_zone_rank(
    record: dict[str, Any],
    admin_area: str,
    tong_ri: str,
    building_name: str,
    address_text: str,
) -> tuple[int, int, int, int, int, int]:
    address_compact = compact_text(address_text)
    building_tokens = tokenize_korean_terms(building_name)
    tong_number = _extract_tong_number(tong_ri)
    zone_text = record["zone_text"]
    zone_compact = compact_text(zone_text)
    zone_token_matches = sum(1 for token in record.get("zone_tokens", []) if token and token in address_compact)
    building_token_matches = sum(1 for token in building_tokens if token and token in zone_compact)
    has_admin_match = int(bool(admin_area and record["admin_area"] == admin_area))
    has_tong_match = int(bool(_tong_or_ri_matches(tong_ri, zone_text)))
    has_tong_text_match = int(bool(tong_ri and _tong_or_ri_matches(tong_ri, zone_text)))
    has_zone_text_in_address = int(bool(zone_compact and zone_compact in address_compact))
    return (
        has_admin_match,
        has_tong_match,
        has_zone_text_in_address,
        building_token_matches,
        zone_token_matches,
        -record.get("row_number", 0),
    )


def choose_school_and_zone(
    context: AppContext,
    admin_area: str,
    tong_ri: str,
    building_name: str,
    address_text: str,
    school_name: str = "",
) -> dict[str, Any]:
    candidates = context.school_zone_records
    if school_name:
        candidates = [record for record in candidates if record["school_name"] == school_name]

    ranked = []
    for record in candidates:
        rank = _rule_zone_rank(record, admin_area, tong_ri, building_name, address_text)
        if rank[:5] != (0, 0, 0, 0, 0):
            ranked.append((rank, record))

    fallback_status = ""
    if not ranked and school_name:
        fallback_status = "학교명불일치후재탐색"
        for record in context.school_zone_records:
            rank = _rule_zone_rank(record, admin_area, tong_ri, building_name, address_text)
            if rank[:5] != (0, 0, 0, 0, 0):
                ranked.append((rank, record))

    if not ranked:
        return {"school_name": school_name, "zone_status": "학교미분류"}

    ranked.sort(key=lambda item: item[0], reverse=True)
    best_rank, best = ranked[0]
    if fallback_status:
        status = fallback_status
    else:
        # 핵심 조건(행정동, 통, 단지 키워드) 2개 이상이면 매칭으로 본다.
        strong_hits = sum(best_rank[:3])
        status = "학교매칭" if strong_hits >= 2 else "학교후보"
    return {**best, "zone_status": status, "zone_score": sum(best_rank[:5])}


def _admin_lookup_from_dataset(address_text: str, building_name: str, context: AppContext) -> tuple[str, str, str, dict[str, Any] | None]:
    normalized = compact_text(address_text)
    building_tokens = tokenize_korean_terms(building_name)
    address_tokens = tokenize_korean_terms(address_text)
    parsed_parcel = parse_jibun_address(address_text)
    unit_info = parse_address_unit_info(address_text)
    address_block_numbers = unit_info["block_numbers"]
    address_unit_numbers = unit_info["unit_numbers"]

    if parsed_parcel["main_number"] is not None:
        parcel_candidates: list[dict[str, Any]] = []
        for record in context.admin_records:
            if not record["legal_area"]:
                continue
            if record["legal_area"] != parsed_parcel["legal_area"]:
                continue
            if parcel_in_ranges(parsed_parcel, record.get("parcel_ranges", [])):
                parcel_candidates.append(record)
        if parcel_candidates:
            if len(parcel_candidates) == 1:
                record = parcel_candidates[0]
                return record["admin_area"], record["tong_ri"], record["ban"], record

            def rank_candidate(record: dict[str, Any]) -> tuple[int, int, int, int, int, int]:
                candidate_tokens = tokenize_korean_terms(" ".join(record.get("building_names", [])))
                def is_meaningful_token(token: str) -> bool:
                    if not token:
                        return False
                    compact = compact_text(token)
                    if len(compact) < 3:
                        return False
                    if re.fullmatch(r"\d+(?:동|호|통|반|단지)?", compact):
                        return False
                    if compact in {"아파트", "단지", "마을", "빌라", "동", "호"}:
                        return False
                    return True
                token_hits = sum(
                    1
                    for token in building_tokens + address_tokens
                    if is_meaningful_token(token) and any(token in candidate for candidate in candidate_tokens)
                )
                has_building_exact = int(bool(building_name and any_contains(" ".join(record.get("building_names", [])), [building_name])))
                has_jurisdiction_text = int(bool(record.get("building_name") and compact_text(record["building_name"]) in normalized))
                record_blocks = [clean_text(value) for value in record.get("block_numbers", [])]
                has_block_match = 0
                has_block_conflict = 0
                if address_block_numbers and record_blocks:
                    if any(block in record_blocks for block in address_block_numbers):
                        has_block_match = 1
                    else:
                        has_block_conflict = 1
                has_unit_match = 0
                if address_unit_numbers and record.get("unit_ranges"):
                    for unit_number in address_unit_numbers:
                        if unit_in_ranges(unit_number, record.get("unit_ranges", []), address_block_numbers[0] if address_block_numbers else ""):
                            has_unit_match = 1
                            break
                has_jurisdiction_full_match = int(bool(record["jurisdiction"] and compact_text(record["jurisdiction"]) in normalized))
                return (
                    has_building_exact,
                    has_block_match,
                    has_unit_match,
                    has_jurisdiction_full_match,
                    token_hits - has_block_conflict,
                    -record["row_number"],
                )

            best = max(parcel_candidates, key=rank_candidate)
            record = best
            return record["admin_area"], record["tong_ri"], record["ban"], record

    # 면 단위 리 주소 fallback:
    # 지번 세부 번지가 관할표에 누락되어도 같은 '리'로 행정동을 식별할 수 있으면
    # 해당 리를 기준으로 학구 매칭이 가능하도록 admin_area/tong_ri를 보정한다.
    legal_area = clean_text(parsed_parcel.get("legal_area", ""))
    if legal_area and legal_area.endswith("리"):
        ri_records = [record for record in context.admin_records if record.get("legal_area") == legal_area]
        if ri_records:
            admin_area = next((r.get("admin_area", "") for r in ri_records if "면" in clean_text(r.get("admin_area", ""))), "")
            if not admin_area:
                admin_area = ri_records[0].get("admin_area", "")
            # 리 이름을 tong_ri로 사용해 용미리/분수리 같은 면단위 학구를 폭넓게 포함한다.
            return admin_area, legal_area, "", None

    if building_tokens or address_tokens:
        for record in context.admin_records:
            candidate_tokens = record.get("building_names", [])
            if candidate_tokens and any_contains(" ".join(candidate_tokens), building_tokens + address_tokens):
                if record["legal_area"] and record["legal_area"] in clean_text(address_text):
                    return record["admin_area"], record["tong_ri"], record["ban"], record

    for record in context.admin_records:
        if record["legal_area"] and record["legal_area"] not in clean_text(address_text):
            continue
        if record["jurisdiction"] and compact_text(record["jurisdiction"]) in normalized:
            return record["admin_area"], record["tong_ri"], record["ban"], record

    dong, main, sub = extract_address_components(clean_text(address_text))
    parcel_key = f"{dong} {main}" + (f"-{sub}" if sub else "") if dong and main else ""
    if parcel_key:
        for record in context.admin_records:
            if parcel_key and compact_text(parcel_key) in compact_text(record["jurisdiction"]):
                return record["admin_area"], record["tong_ri"], record["ban"], record

    return "", "", "", None


def _fallback_admin_lookup(jibun_address: str, context: AppContext) -> tuple[str, str]:
    dong_name, main_number, sub_number = extract_address_components(jibun_address)
    row = pd.Series(
        {
            "지번주소": jibun_address,
            "std_address": standardize_address(jibun_address),
            "building_name": extract_building_name(jibun_address),
            "dong_name": dong_name,
            "main_number": main_number,
            "sub_number": sub_number,
        }
    )
    mapped_dong, mapped_tong = map_address_to_school_area(
        row,
        manual_mapping,
        {},
        {},
        context.fallback_district_ranges,
    )
    return mapped_dong, mapped_tong


def classify_student(student: dict[str, Any], context: AppContext, school_name: str = "") -> dict[str, Any]:
    road_address = student["road_address"]
    postal_code = student.get("postal_code", "")
    requested_school_name = school_name
    source_jibun = student["source_jibun_address"]
    base_address = source_jibun or road_address
    is_jibun = bool(source_jibun) or bool(re.search(r"[가-힣]+동\s*\d+", base_address))

    jibun_address = source_jibun
    if not jibun_address and road_address:
        normalized_road = _normalize_road_for_jibun_lookup(road_address)
        jibun_address = get_jibeon_address(normalized_road or road_address, expected_zip=postal_code)

    match_text = jibun_address or road_address
    building_name = extract_building_name(match_text)
    anchor_rule = _resolve_address_anchor(
        road_address=road_address,
        jibun_address=jibun_address,
        building_name=building_name or "",
        requested_school_name=requested_school_name,
    )
    if anchor_rule:
        admin_area = clean_text(anchor_rule.get("admin_area", ""))
        tong_ri = clean_text(anchor_rule.get("tong_ri", ""))
        ban = ""
        admin_record = None
    else:
        admin_area, tong_ri, ban, admin_record = _admin_lookup_from_dataset(match_text, building_name, context)
        if (not admin_area or not tong_ri) and jibun_address:
            fallback_dong, fallback_tong = _fallback_admin_lookup(jibun_address, context)
            admin_area = admin_area or fallback_dong
            tong_ri = tong_ri or fallback_tong

    school_zone_match = choose_school_and_zone(
        context=context,
        admin_area=admin_area,
        tong_ri=tong_ri,
        building_name=building_name,
        address_text=" ".join(filter(None, [road_address, jibun_address])),
        school_name="",
    )
    resolved_school_name = school_zone_match.get("school_name", "") or requested_school_name

    report_match = match_report_row(
        context=context,
        admin_area=admin_area,
        tong_ri=tong_ri,
        building_name=building_name,
        address_text=" ".join(filter(None, [road_address, jibun_address])),
        school_name="",
        admin_building_names=admin_record.get("building_names", []) if admin_record else [],
    )
    actual_school_name = report_match.get("school_name", "") or resolved_school_name
    if anchor_rule and anchor_rule.get("school_name"):
        anchor_school = clean_text(anchor_rule["school_name"])
        forced_report_match = match_report_row(
            context=context,
            admin_area=admin_area,
            tong_ri=tong_ri,
            building_name=building_name,
            address_text=" ".join(filter(None, [road_address, jibun_address])),
            school_name=anchor_school,
            admin_building_names=admin_record.get("building_names", []) if admin_record else [],
        )
        if forced_report_match.get("row_number"):
            report_match = forced_report_match
        actual_school_name = anchor_school
    if actual_school_name:
        refined_report_match = match_report_row(
            context=context,
            admin_area=admin_area,
            tong_ri=tong_ri,
            building_name=building_name,
            address_text=" ".join(filter(None, [road_address, jibun_address])),
            school_name=actual_school_name,
            admin_building_names=admin_record.get("building_names", []) if admin_record else [],
        )
        if refined_report_match.get("row_number"):
            report_match = refined_report_match
            actual_school_name = report_match.get("school_name", "") or actual_school_name
    requested_report_match = {}
    if requested_school_name:
        requested_report_match = match_report_row(
            context=context,
            admin_area=admin_area,
            tong_ri=tong_ri,
            building_name=building_name,
            address_text=" ".join(filter(None, [road_address, jibun_address])),
            school_name=requested_school_name,
            admin_building_names=admin_record.get("building_names", []) if admin_record else [],
        )
        if (
            not requested_report_match.get("row_number")
            or not requested_report_match.get("status", "").startswith("정확한매칭")
        ):
            token_based_match = _match_report_row_by_address_tokens(
                context=context,
                requested_school_name=requested_school_name,
                address_text=" ".join(filter(None, [road_address, jibun_address])),
            )
            if token_based_match.get("row_number"):
                requested_report_match = token_based_match
        if (
            requested_school_name
            and actual_school_name
            and requested_school_name != actual_school_name
            and _is_joint_zone_pair(requested_report_match, report_match, requested_school_name, actual_school_name)
        ):
            report_match = requested_report_match
            actual_school_name = requested_school_name
        elif (
            requested_report_match.get("row_number")
            and requested_school_name
            and requested_school_name != actual_school_name
            and requested_report_match.get("status", "").startswith("정확한매칭")
        ):
            report_match = requested_report_match
            actual_school_name = requested_school_name
        elif (
            requested_report_match.get("row_number")
            and requested_school_name
            and requested_school_name != actual_school_name
            and _tong_or_ri_matches(tong_ri, requested_report_match.get("tong_ri", ""))
        ):
            # 면/리 단위 주소는 번지 누락이 잦으므로, 요청학교의 리 매칭이 성립하면 요청학교를 우선한다.
            report_match = requested_report_match
            actual_school_name = requested_school_name
    school_alignment_status = "일치"
    zone_match_status = school_zone_match.get("zone_status", "학교미분류")
    match_status = report_match.get("status", "미분류")
    match_reason = report_match.get("reason", "")
    outside_paju = _is_outside_paju_address(road_address, jibun_address)

    if outside_paju:
        special_row = _find_outside_special_row(context, requested_school_name or actual_school_name)
        if special_row:
            report_match = {**special_row, "status": "관외(주소확인요망)"}
            if not actual_school_name:
                actual_school_name = special_row.get("school_name", "") or requested_school_name
        else:
            report_match = {**report_match, "status": "관외(주소확인요망)"}
        school_alignment_status = "관외(주소확인요망)"
        zone_match_status = "관외(주소확인요망)"
        match_status = "관외(주소확인요망)"
        match_reason = "경기도 파주시 외 주소"
    elif requested_school_name and actual_school_name and requested_school_name != actual_school_name:
        school_alignment_status = "관내(학구위반)"
        zone_match_status = "관내(학구위반)"
        match_status = "관내(학구위반)"
        match_reason = "요청학교와 판정학교가 다름"

    return {
        "grade": student["grade"],
        "homeroom_ban": clean_text(student.get("homeroom_ban", "")),
        "attendance_no": clean_text(student.get("attendance_no", "")),
        "name": student["name"],
        "input_type": "jibun" if is_jibun else "road",
        "requested_school_name": requested_school_name,
        "postal_code": postal_code,
        "road_address": road_address,
        "jibun_address": jibun_address,
        "building_name": building_name,
        "admin_area": admin_area,
        "tong_ri": tong_ri,
        "ban": ban,
        "admin_building_names": admin_record.get("building_names", []) if admin_record else [],
        "school_alignment_status": school_alignment_status,
        "school_name": actual_school_name,
        "sheet_name": report_match.get("sheet_name", "재학생 현황(관내)"),
        "report_row_number": report_match.get("row_number"),
        "report_admin_area": report_match.get("admin_area", ""),
        "report_tong_ri": report_match.get("tong_ri", ""),
        "report_extra_area": report_match.get("extra_area", ""),
        "zone_match_status": zone_match_status,
        "zone_match_score": school_zone_match.get("zone_score", 0),
        "match_status": match_status,
        "match_reason": match_reason,
    }


def match_report_row(
    context: AppContext,
    admin_area: str,
    tong_ri: str,
    building_name: str,
    address_text: str,
    school_name: str = "",
    admin_building_names: list[str] | None = None,
) -> dict[str, Any]:
    candidate_rows = [
        record
        for record in context.report_template_records
        if record["sheet_name"] == "재학생 현황(관내)" and record["category"] == "normal"
    ]

    if school_name:
        candidate_rows = [record for record in candidate_rows if record["school_name"] == school_name]

    if admin_area:
        candidate_rows = [record for record in candidate_rows if record["admin_area"] == admin_area]

    if tong_ri:
        filtered = []
        for record in candidate_rows:
            if _tong_or_ri_matches(tong_ri, record["tong_ri"]):
                filtered.append(record)
        if filtered:
            candidate_rows = filtered

    if not candidate_rows:
        return {"status": "미분류"}

    if len(candidate_rows) > 1:
        exclude_rows = [row for row in candidate_rows if "제외" in row["extra_area"]]
        explicit_rows = [row for row in candidate_rows if row["extra_area"] and "제외" not in row["extra_area"]]
        combined_tokens = tokenize_korean_terms(building_name) + tokenize_korean_terms(" ".join(admin_building_names or []))
        explicit_match = []
        for row in explicit_rows:
            row_tokens = tokenize_korean_terms(row["extra_area"])
            if any(token and any(token in candidate for candidate in row_tokens) for token in combined_tokens):
                explicit_match.append(row)
        if exclude_rows and explicit_rows and not explicit_match:
            candidate_rows = exclude_rows
        elif explicit_match:
            candidate_rows = explicit_match

    if len(candidate_rows) == 1:
        row = candidate_rows[0]
        return {**row, "status": "정확한매칭", "reason": "후보가 1개로 확정"}

    building_tokens = tokenize_korean_terms(building_name)
    address_compact = compact_text(address_text)
    ranked: list[tuple[tuple[int, int, int, int], dict[str, Any]]] = []
    for record in candidate_rows:
        extra_compact = compact_text(record.get("extra_area", ""))
        token_hit_count = sum(1 for token in record.get("extra_tokens", []) if token and token in address_compact)
        building_hit_count = sum(1 for token in building_tokens if token and token in extra_compact)
        has_extra_full_match = int(bool(extra_compact and extra_compact in address_compact))
        has_tong_text_match = int(bool(record["tong_ri"] and compact_text(record["tong_ri"]) in address_compact))
        rank = (has_extra_full_match, building_hit_count, token_hit_count, has_tong_text_match)
        ranked.append((rank, record))

    ranked.sort(key=lambda item: item[0], reverse=True)
    best_rank, best_record = ranked[0]
    if best_rank[0] == 1 or (best_rank[1] + best_rank[2]) >= 2:
        return {**best_record, "status": "정확한매칭", "reason": "행정구역/단지 토큰 규칙 일치"}
    if any(best_rank):
        return {**best_record, "status": "일부매칭", "reason": "토큰 일부 일치(후보 중 최선)"}

    return {**best_record, "status": "일부매칭", "reason": "명확 근거 부족, 후보 기본값"}


def _match_report_row_by_address_tokens(
    context: AppContext,
    requested_school_name: str,
    address_text: str,
) -> dict[str, Any]:
    if not requested_school_name:
        return {}
    address_compact = compact_text(address_text)
    if not address_compact:
        return {}

    candidate_rows = [
        record
        for record in context.report_template_records
        if record["sheet_name"] == "재학생 현황(관내)"
        and record["category"] == "normal"
        and record["school_name"] == requested_school_name
        and record.get("extra_area")
    ]
    if not candidate_rows:
        return {}

    scored: list[tuple[int, dict[str, Any]]] = []
    for row in candidate_rows:
        score = 0
        extra_compact = compact_text(row.get("extra_area", ""))
        if extra_compact and extra_compact in address_compact:
            score += 120
        for token in row.get("extra_tokens", []):
            token_compact = compact_text(token)
            if token_compact and token_compact in address_compact:
                score += 45
        if score > 0:
            scored.append((score, row))

    if not scored:
        return {}
    scored.sort(key=lambda item: item[0], reverse=True)
    if scored[0][0] < 90:
        return {}
    return {**scored[0][1], "status": "정확한매칭(주소토큰)", "reason": "요청학교 행의 주소 토큰 직접 일치"}


def _student_list_dataframe(students: list[dict[str, Any]]) -> pd.DataFrame:
    """
    모든 산출물에서 공통으로 사용할 학생 목록 DataFrame.
    학년/반/번호를 항상 앞에 고정해 누락되지 않게 한다.
    """
    rows = [
        {
            "학년": row.get("grade", ""),
            "반": clean_text(row.get("homeroom_ban", "")),
            "번호": clean_text(row.get("attendance_no", "")),
            "성명": row.get("name", ""),
            "우편번호": clean_text(row.get("postal_code", "")),
            "도로명주소": clean_text(row.get("road_address", "")),
            "지번주소": clean_text(row.get("jibun_address", "")),
            "행정동": clean_text(row.get("admin_area", "")),
            "통리": clean_text(row.get("tong_ri", "")),
            "학교명": clean_text(row.get("school_name", "")),
            "학교매칭": clean_text(row.get("zone_match_status", "")),
            "행매칭": clean_text(row.get("match_status", "")),
            "매칭사유": clean_text(row.get("match_reason", "")),
        }
        for row in students
    ]
    return pd.DataFrame(rows)


def build_report(processed_students: list[dict[str, Any]], school_name: str = "", base_date: str = "") -> dict[str, Any]:
    context = get_context()
    config = load_admin_config()
    display_base_date = _format_base_date(base_date)
    template_path = _resolve_report_template_path(config.data_sources.report_template_source)
    aggregate: dict[tuple[str, int], dict[str, Any]] = {}
    report_records = context.report_template_records
    for record in report_records:
        aggregate[(record["sheet_name"], record["row_number"])] = {
            **record,
            **{column: 0 for column in GRADE_COLUMNS},
            "계": 0,
        }

    unmatched = []
    for student in processed_students:
        row_number = student.get("report_row_number")
        sheet_name = student.get("sheet_name")
        grade = student["grade"]

        if not row_number or (sheet_name, row_number) not in aggregate or grade not in range(1, 7):
            unmatched.append(student)
            continue

        grade_col = f"{grade}학년"
        aggregate[(sheet_name, row_number)][grade_col] += 1
        aggregate[(sheet_name, row_number)]["계"] += 1

    output = BytesIO()
    all_students_df = _student_list_dataframe(processed_students)
    if template_path and os.path.exists(template_path):
        workbook = load_workbook(template_path)
        if display_base_date:
            for sheet_name in workbook.sheetnames:
                ws = workbook[sheet_name]
                if sheet_name == "재학생 현황(관내)":
                    ws["L3"] = f"기준일: {display_base_date}"
                elif sheet_name == "재학생 현황(관외)":
                    ws["M3"] = f"기준일: {display_base_date}"

        for (sheet_name, row_number), row_data in aggregate.items():
            ws = workbook[sheet_name]
            if sheet_name == "재학생 현황(관내)":
                start_col = 5
                total_col = 11
            else:
                start_col = 6
                total_col = 12
            for offset, column in enumerate(GRADE_COLUMNS):
                ws.cell(row=row_number, column=start_col + offset, value=row_data[column])
            ws.cell(row=row_number, column=total_col, value=row_data["계"])
        if "학생처리목록" in workbook.sheetnames:
            del workbook["학생처리목록"]
        ws_students = workbook.create_sheet("학생처리목록")
        ws_students.append(list(all_students_df.columns))
        for row in all_students_df.itertuples(index=False, name=None):
            ws_students.append(list(row))
        workbook.save(output)
    else:
        fallback_rows = list(aggregate.values())
        fallback_rows.sort(key=lambda row: (row["sheet_name"], row["row_number"]))
        fallback_df = pd.DataFrame(fallback_rows)
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            fallback_df.to_excel(writer, sheet_name="집계결과", index=False)
            all_students_df.to_excel(writer, sheet_name="학생처리목록", index=False)
            if display_base_date:
                pd.DataFrame([{"기준일": display_base_date}]).to_excel(writer, sheet_name="메타", index=False)
    output.seek(0)

    table_rows = [
        row for row in aggregate.values()
        if row["category"] == "normal" and row["계"] > 0
    ]
    table_rows.sort(key=lambda row: (row["school_name"], row["admin_area"], row["tong_ri"], row["extra_area"]))

    details_df = all_students_df.copy()
    details_output = BytesIO()
    with pd.ExcelWriter(details_output, engine="openpyxl") as writer:
        details_df.to_excel(writer, sheet_name="students", index=False)
    details_output.seek(0)

    issue_students = [
        row for row in processed_students
        if row.get("match_status") in {"관내(학구위반)", "관외(주소확인요망)", "미분류"}
    ]
    issue_df = pd.DataFrame(issue_students)
    if not issue_df.empty:
        issue_rename = {
            "grade": "학년",
            "homeroom_ban": "반",
            "attendance_no": "번호",
            "name": "성명",
        }
        issue_df = issue_df.rename(columns={k: v for k, v in issue_rename.items() if k in issue_df.columns})
        issue_front = ["학년", "반", "번호", "성명"]
        issue_ordered = [c for c in issue_front if c in issue_df.columns]
        issue_ordered.extend([c for c in issue_df.columns if c not in issue_ordered])
        issue_df = issue_df[issue_ordered]
    issue_output = BytesIO()
    with pd.ExcelWriter(issue_output, engine="openpyxl") as writer:
        issue_df.to_excel(writer, sheet_name="issues", index=False)
    issue_output.seek(0)

    summary = {
        "total_students": len(processed_students),
        "matched_students": len(processed_students) - len(unmatched),
        "unmatched_students": len(unmatched),
        "exact_matches": sum(1 for row in processed_students if str(row["match_status"]).startswith("정확한매칭")),
        "candidate_matches": sum(1 for row in processed_students if row["match_status"] == "일부매칭"),
        "school_mismatches": sum(1 for row in processed_students if row["match_status"] == "관내(학구위반)"),
        "outside_address_checks": sum(1 for row in processed_students if row["match_status"] == "관외(주소확인요망)"),
        "schools": sorted({row["school_name"] for row in table_rows if row["school_name"]}),
    }

    requested_school_workbook = b""
    if school_name:
        requested_records = [
            record for record in context.report_template_records
            if record["school_name"] == school_name
        ]
        requested_aggregate: dict[tuple[str, int], dict[str, Any]] = {}
        for record in requested_records:
            requested_aggregate[(record["sheet_name"], record["row_number"])] = {
                **record,
                **{column: 0 for column in GRADE_COLUMNS},
                "계": 0,
            }

        special_row = next(
            (
                record for record in requested_records
                if record["sheet_name"] == "재학생 현황(관내)" and record["category"] == "special" and record["admin_area"] == "기타(통학구역외)"
            ),
            None,
        )
        summary_row = next(
            (
                record for record in requested_records
                if record["sheet_name"] == "재학생 현황(관내)" and record["category"] == "summary" and record["admin_area"] == "합계"
            ),
            None,
        )

        for student in processed_students:
            grade = student["grade"]
            if grade not in range(1, 7):
                continue
            grade_col = f"{grade}학년"

            if student.get("match_status") in {"관내(학구위반)", "관외(주소확인요망)", "미분류"}:
                if special_row:
                    bucket = requested_aggregate[(special_row["sheet_name"], special_row["row_number"])]
                    bucket[grade_col] += 1
                    bucket["계"] += 1
                continue

            if student.get("school_name") != school_name:
                continue
            row_number = student.get("report_row_number")
            sheet_name = student.get("sheet_name")
            if not row_number or (sheet_name, row_number) not in requested_aggregate:
                continue
            bucket = requested_aggregate[(sheet_name, row_number)]
            bucket[grade_col] += 1
            bucket["계"] += 1

        if summary_row:
            summary_bucket = requested_aggregate[(summary_row["sheet_name"], summary_row["row_number"])]
            normal_and_special_rows = [
                row for row in requested_aggregate.values()
                if row["sheet_name"] == summary_row["sheet_name"] and row["category"] in {"normal", "special"}
            ]
            for grade_col in GRADE_COLUMNS:
                summary_bucket[grade_col] = sum(row[grade_col] for row in normal_and_special_rows)
            summary_bucket["계"] = sum(row["계"] for row in normal_and_special_rows)

        requested_output = BytesIO()
        requested_students = [row for row in processed_students if row.get("school_name") == school_name]
        requested_students_df = _student_list_dataframe(requested_students)
        if template_path and os.path.exists(template_path):
            requested_workbook = load_workbook(template_path)
            if display_base_date:
                for sheet_name in requested_workbook.sheetnames:
                    ws = requested_workbook[sheet_name]
                    if sheet_name == "재학생 현황(관내)":
                        ws["L3"] = f"기준일: {display_base_date}"
                    elif sheet_name == "재학생 현황(관외)":
                        ws["M3"] = f"기준일: {display_base_date}"

            for (sheet_name, row_number), row_data in requested_aggregate.items():
                ws = requested_workbook[sheet_name]
                if sheet_name == "재학생 현황(관내)":
                    start_col = 5
                    total_col = 11
                else:
                    start_col = 6
                    total_col = 12
                for offset, column in enumerate(GRADE_COLUMNS):
                    ws.cell(row=row_number, column=start_col + offset, value=row_data[column])
                ws.cell(row=row_number, column=total_col, value=row_data["계"])
            if "학생처리목록" in requested_workbook.sheetnames:
                del requested_workbook["학생처리목록"]
            ws_students = requested_workbook.create_sheet("학생처리목록")
            ws_students.append(list(requested_students_df.columns))
            for row in requested_students_df.itertuples(index=False, name=None):
                ws_students.append(list(row))
            requested_workbook.save(requested_output)
        else:
            requested_rows = list(requested_aggregate.values())
            requested_rows.sort(key=lambda row: (row["sheet_name"], row["row_number"]))
            requested_df = pd.DataFrame(requested_rows)
            with pd.ExcelWriter(requested_output, engine="openpyxl") as writer:
                requested_df.to_excel(writer, sheet_name="우리학교집계", index=False)
                requested_students_df.to_excel(writer, sheet_name="학생처리목록", index=False)
                if display_base_date:
                    pd.DataFrame([{"기준일": display_base_date}]).to_excel(writer, sheet_name="메타", index=False)
        requested_output.seek(0)
        requested_school_workbook = requested_output.getvalue()

    return {
        "summary": summary,
        "table_rows": table_rows,
        "unmatched_students": unmatched[:100],
        "report_workbook": output.getvalue(),
        "requested_school_workbook": requested_school_workbook,
        "detail_workbook": details_output.getvalue(),
        "issue_workbook": issue_output.getvalue(),
    }


RESULT_STORE: dict[str, dict[str, bytes]] = {}
RESULT_DATA_STORE: dict[str, dict[str, Any]] = {}
JOB_STORE: dict[str, dict[str, Any]] = {}
JOB_LOCK = threading.Lock()


def _now_ts() -> float:
    return time.time()


def _purge_expired_data() -> None:
    config = load_admin_config()
    now = _now_ts()
    result_ttl = max(60, config.runtime.result_ttl_minutes * 60)
    job_ttl = max(60, config.runtime.job_ttl_minutes * 60)

    with JOB_LOCK:
        expired_results = [
            result_id for result_id, payload in RESULT_DATA_STORE.items()
            if now - payload.get("created_at", now) > result_ttl
        ]
        for result_id in expired_results:
            RESULT_DATA_STORE.pop(result_id, None)
            RESULT_STORE.pop(result_id, None)

        expired_jobs = [
            job_id for job_id, payload in JOB_STORE.items()
            if now - payload.get("created_at", now) > job_ttl
        ]
        for job_id in expired_jobs:
            JOB_STORE.pop(job_id, None)


def start_cleanup_worker() -> None:
    global _CLEANUP_THREAD_STARTED
    if _CLEANUP_THREAD_STARTED:
        return
    _CLEANUP_THREAD_STARTED = True

    def worker() -> None:
        while True:
            try:
                _purge_expired_data()
                sleep_seconds = max(30, load_admin_config().runtime.cleanup_interval_seconds)
            except Exception:
                sleep_seconds = 60
            time.sleep(sleep_seconds)

    thread = threading.Thread(target=worker, name="cleanup-worker", daemon=True)
    thread.start()


def _set_job_state(job_id: str, **updates) -> None:
    with JOB_LOCK:
        if job_id in JOB_STORE:
            JOB_STORE[job_id].update(updates)


def build_jibun_export_workbook(students: list[dict[str, Any]]) -> bytes:
    """도로명→지번 변환 결과만 담은 엑셀 바이너리를 만든다."""
    rows: list[dict[str, Any]] = []
    for student in students:
        rows.append(
            {
                "학년": student.get("grade", ""),
                "반": clean_text(student.get("homeroom_ban", "")),
                "번호": clean_text(student.get("attendance_no", "")),
                "성명": student.get("name", ""),
                "우편번호": clean_text(student.get("postal_code", "")),
                "도로명주소": clean_text(student.get("road_address", "")),
                "지번주소": clean_text(student.get("source_jibun_address", "")),
            }
        )
    buffer = BytesIO()
    pd.DataFrame(rows).to_excel(buffer, sheet_name="지번변환", index=False)
    buffer.seek(0)
    return buffer.getvalue()


def process_address_conversion_request(
    files: list[tuple[str, bytes]],
    progress_callback=None,
) -> dict[str, Any]:
    """
    통학구역 매칭 없이 명렬표의 도로명 주소만 행정안전부 API로 지번 주소로 채운다.
    """
    if progress_callback:
        progress_callback(stage="학생 파일 읽는 중", current=0, total=0)
    students = load_student_files(files)
    _prefill_jibun_addresses(students, progress_callback=progress_callback)
    if progress_callback:
        progress_callback(stage="엑셀 생성 중", current=len(students), total=len(students))
    workbook_bytes = build_jibun_export_workbook(students)
    result_id = uuid4().hex
    created_at = _now_ts()
    with_jibun = sum(1 for row in students if clean_text(row.get("source_jibun_address", "")))
    summary = {
        "total_students": len(students),
        "with_jibun": with_jibun,
        "missing_jibun": len(students) - with_jibun,
    }
    RESULT_STORE[result_id] = {"converted": workbook_bytes}
    RESULT_DATA_STORE[result_id] = {
        "created_at": created_at,
        "mode": "address_convert",
        "summary": summary,
        "result_id": result_id,
    }
    if progress_callback:
        progress_callback(stage="완료", current=len(students), total=len(students))
    return RESULT_DATA_STORE[result_id]


def process_request(
    files: list[tuple[str, bytes]],
    school_name: str = "",
    base_date: str = "",
    progress_callback=None,
) -> dict[str, Any]:
    context = get_context()
    if progress_callback:
        progress_callback(stage="학생 파일 읽는 중", current=0, total=0)
    students = load_student_files(files)
    _prefill_jibun_addresses(students, progress_callback=progress_callback)
    processed = []
    total = len(students)
    for index, student in enumerate(students, start=1):
        processed.append(classify_student(student, context, school_name=school_name))
        if progress_callback:
            progress_callback(stage="주소 변환 및 매칭 중", current=index, total=total)
    if progress_callback:
        progress_callback(stage="서식 생성 중", current=total, total=total)
    result = build_report(processed, school_name=school_name, base_date=base_date)
    result_id = uuid4().hex
    created_at = _now_ts()
    RESULT_STORE[result_id] = {
        "report": result["report_workbook"],
        "school_report": result["requested_school_workbook"],
        "details": result["detail_workbook"],
        "issues": result["issue_workbook"],
    }
    result["created_at"] = created_at
    RESULT_DATA_STORE[result_id] = result
    result["result_id"] = result_id
    if progress_callback:
        progress_callback(stage="완료", current=total, total=total)
    return result


def _prefill_jibun_addresses(students: list[dict[str, Any]], progress_callback=None) -> None:
    # 같은 도로명 주소가 반복되는 경우 외부 API 호출을 1회로 줄인다.
    unique_targets: dict[tuple[str, str], None] = {}
    for student in students:
        if student.get("source_jibun_address"):
            continue
        road_address = clean_text(student.get("road_address", ""))
        if not road_address:
            continue
        postal_code = clean_text(student.get("postal_code", ""))
        unique_targets[(road_address, postal_code)] = None

    if not unique_targets:
        return

    conversion_cache: dict[tuple[str, str], str] = {}
    total_unique = len(unique_targets)
    for index, (road_address, postal_code) in enumerate(unique_targets.keys(), start=1):
        normalized_road = _normalize_road_for_jibun_lookup(road_address)
        converted = get_jibeon_address(normalized_road or road_address, expected_zip=postal_code) or ""
        conversion_cache[(road_address, postal_code)] = converted
        if progress_callback:
            progress_callback(stage="중복 주소 통합 변환 중", current=index, total=total_unique)

    for student in students:
        if student.get("source_jibun_address"):
            continue
        road_address = clean_text(student.get("road_address", ""))
        postal_code = clean_text(student.get("postal_code", ""))
        cached = conversion_cache.get((road_address, postal_code), "")
        if cached:
            student["source_jibun_address"] = cached


def get_download(result_id: str, kind: str) -> bytes | None:
    _purge_expired_data()
    bucket = RESULT_STORE.get(result_id, {})
    return bucket.get(kind)


def get_result_data(result_id: str) -> dict[str, Any] | None:
    _purge_expired_data()
    return RESULT_DATA_STORE.get(result_id)


def start_job(files: list[tuple[str, bytes]], school_name: str = "", base_date: str = "") -> str:
    job_id = uuid4().hex
    created_at = _now_ts()
    with JOB_LOCK:
        JOB_STORE[job_id] = {
            "created_at": created_at,
            "status": "queued",
            "stage": "대기 중",
            "current": 0,
            "total": 0,
            "error": "",
            "result_id": "",
        }

    def progress_callback(stage: str, current: int, total: int) -> None:
        _set_job_state(
            job_id,
            status="running" if stage != "완료" else "completed",
            stage=stage,
            current=current,
            total=total,
        )

    def worker() -> None:
        try:
            result = process_request(
                files,
                school_name=school_name,
                base_date=base_date,
                progress_callback=progress_callback,
            )
            _set_job_state(
                job_id,
                status="completed",
                stage="완료",
                current=JOB_STORE[job_id]["current"],
                total=JOB_STORE[job_id]["total"],
                result_id=result["result_id"],
            )
        except Exception as exc:
            _set_job_state(job_id, status="failed", stage="실패", error=str(exc))

    thread = threading.Thread(target=worker, name=f"job-{job_id[:8]}", daemon=True)
    thread.start()
    return job_id


def start_job_address_conversion(files: list[tuple[str, bytes]]) -> str:
    job_id = uuid4().hex
    created_at = _now_ts()
    with JOB_LOCK:
        JOB_STORE[job_id] = {
            "created_at": created_at,
            "status": "queued",
            "stage": "대기 중",
            "current": 0,
            "total": 0,
            "error": "",
            "result_id": "",
        }

    def progress_callback(stage: str, current: int, total: int) -> None:
        _set_job_state(
            job_id,
            status="running" if stage != "완료" else "completed",
            stage=stage,
            current=current,
            total=total,
        )

    def worker() -> None:
        try:
            result = process_address_conversion_request(files, progress_callback=progress_callback)
            _set_job_state(
                job_id,
                status="completed",
                stage="완료",
                current=JOB_STORE[job_id]["current"],
                total=JOB_STORE[job_id]["total"],
                result_id=result["result_id"],
            )
        except Exception as exc:
            _set_job_state(job_id, status="failed", stage="실패", error=str(exc))

    thread = threading.Thread(target=worker, name=f"addr-job-{job_id[:8]}", daemon=True)
    thread.start()
    return job_id


def get_job(job_id: str) -> dict[str, Any] | None:
    _purge_expired_data()
    with JOB_LOCK:
        job = JOB_STORE.get(job_id)
        return dict(job) if job else None
